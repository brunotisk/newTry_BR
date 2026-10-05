from decimal import Decimal
from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from servicos import vendas_import as vendas


def test_normalizar_codigo_interno_recoloca_zeros():
    assert vendas.normalizar_codigo_interno(24727) == "0024727"
    assert vendas.normalizar_codigo_interno("0024727") == "0024727"
    assert vendas.normalizar_codigo_interno("24727.0") == "0024727"
    assert vendas.normalizar_codigo_interno("ABC-1") == "ABC-1"


def test_parse_moeda_aceita_formatos_brasileiros():
    assert vendas._parse_moeda("R$ 1.234,56") == Decimal("1234.56")
    assert vendas._parse_moeda("79,90") == Decimal("79.90")
    assert vendas._parse_moeda(12.5) == Decimal("12.5")
    assert vendas._parse_moeda("abc") == Decimal("0")


def test_parse_data_normaliza_excel_e_texto():
    from datetime import datetime
    assert vendas._parse_data(datetime(2026, 10, 5, 12, 30)) == "2026-10-05"
    assert vendas._parse_data("05/10/2026") == "2026-10-05"
    assert vendas._parse_data("2026-10-05") == "2026-10-05"
    with pytest.raises(ValueError):
        vendas._parse_data("05-10-2026")


def test_calcular_comissao_sem_parceiro_retorna_none(sb):
    assert vendas.calcular_comissao_venda(sb, 10, 100) == (None, None)


def test_calcular_comissao_usa_parceiro_ativo(sb):
    sb.responses[("parceiros", "select")] = [{"percentual_comissao": 12.5}]
    assert vendas.calcular_comissao_venda(sb, 10, 80) == (12.5, 10.0)


def test_obter_ou_criar_lista_id_reutiliza_cadastro_existente(sb):
    sb.responses[("canais_venda", "select")] = [{"id": 4}]
    assert vendas.obter_ou_criar_lista_id(sb, "canais_venda", "Feira") == 4
    assert not sb.calls_for("canais_venda", "insert")


def test_obter_ou_criar_lista_id_cria_cadastro_ausente(sb):
    sb.responses[("canais_venda", "select")] = []
    sb.responses[("canais_venda", "insert")] = [{"id": 5}]
    assert vendas.obter_ou_criar_lista_id(sb, "canais_venda", "Loja") == 5
    assert sb.last_call("canais_venda", "insert").payload == {"nome": "Loja"}


def test_inserir_venda_baixa_estoque_e_grava_saida(sb, monkeypatch):
    sb.responses[("vendas", "insert")] = [{"id": 30}]
    sb.responses[("estoque", "select")] = [{"quantidade_atual": 10}]
    sb.responses[("estoque", "upsert")] = [{"id": 1}]
    sb.responses[("estoque_movimentos", "insert")] = [{"id": 31}]

    linha = {
        "canal_venda_id": 1, "status_id": 2, "forma_pagamento_id": 3,
        "detalhe_feira_id": None, "quantidade": Decimal("2"),
        "valor_lista": Decimal("100"), "valor_desconto": Decimal("10"),
        "valor_final": Decimal("90"), "data_venda": "2026-10-05",
        "cliente": "Maria", "observacao": "teste",
    }
    venda_id = vendas.inserir_venda_e_baixar_estoque(sb, linha, 7, motivo_saida="Importação XLSX")
    assert venda_id == 30
    payload = sb.last_call("vendas", "insert").payload
    assert payload["produto_id"] == 7
    assert payload["valor_lista"] == 100.0
    assert payload["valor_final"] == 90.0
    movimento = sb.last_call("estoque_movimentos", "insert").payload
    assert movimento["tipo"] == "saida"
    assert movimento["quantidade"] == -2.0
    assert movimento["venda_id"] == 30
    assert movimento["motivo"] == "Importação XLSX"


def _dados_validacao(sb):
    sb.responses[("canais_venda", "select")] = [{"id": 1, "nome": "Loja"}, {"id": 2, "nome": "Feira"}]
    sb.responses[("formas_pagamento", "select")] = [{"id": 3, "descricao": "Pix"}]
    sb.responses[("status_venda", "select")] = [{"id": 4, "nome": "Pago"}]
    sb.responses[("detalhes_feira", "select")] = [{"id": 5, "nome_feira": "Feira Centro"}]
    sb.responses[("clientes", "select")] = [{"id": 6, "nome": "Cliente Antigo"}]
    sb.responses[("produtos", "select")] = [{"id": 7, "codigo_interno": "0024727", "descricao": "Brinco", "estoque": 5}]
    sb.responses[("estoque", "select")] = [{"produto_id": 7, "quantidade_atual": 5}]
    sb.responses[("vendas", "select")] = []


def _planilha(rows, tmp_path):
    path = tmp_path / "vendas.xlsx"
    df = pd.DataFrame(rows)
    df.to_excel(path, index=False)
    return str(path)


def test_validar_planilha_classifica_produto_inexistente_sem_gravar(sb, tmp_path):
    _dados_validacao(sb)
    path = _planilha([{
        "Canal": "Loja", "Código Interno": "99999", "Quantidade": 1,
        "Valor Lista": 100, "Desconto": 0, "Valor Final": 100,
        "Status": "Pago", "Data": "05/10/2026", "Cliente": "Novo",
        "Forma Pagamento": "Pix",
    }], tmp_path)
    resultado = vendas.validar_planilha_vendas(path, sb)
    linha = resultado["linhas"][0]
    assert linha["status_linha"] == "PRODUTO_NAO_CADASTRADO"
    assert resultado["produtos_nao_cadastrados"] == 1
    assert not sb.calls_for("vendas", "insert")


def test_validar_planilha_detecta_estoque_insuficiente(sb, tmp_path):
    _dados_validacao(sb)
    path = _planilha([{
        "Canal": "Loja", "Código Interno": "24727", "Quantidade": 6,
        "Valor Lista": 100, "Desconto": 0, "Valor Final": 100,
        "Status": "Pago", "Data": "05/10/2026", "Cliente": "Novo",
        "Forma Pagamento": "Pix",
    }], tmp_path)
    resultado = vendas.validar_planilha_vendas(path, sb)
    linha = resultado["linhas"][0]
    assert linha["status_linha"] == "ESTOQUE_INSUFICIENTE"
    assert linha["estoque_insuficiente"] is True
    assert resultado["estoque_insuficiente"] == 1


def test_validar_planilha_cliente_novo_e_permitido(sb, tmp_path):
    _dados_validacao(sb)
    path = _planilha([{
        "Canal": "Loja", "Código Interno": "24727", "Quantidade": 1,
        "Valor Lista": 100, "Desconto": 0, "Valor Final": 100,
        "Status": "Pago", "Data": "05/10/2026", "Cliente": "Cliente Novo",
        "Forma Pagamento": "Pix",
    }], tmp_path)
    resultado = vendas.validar_planilha_vendas(path, sb)
    linha = resultado["linhas"][0]
    assert linha["status_linha"] == "CLIENTE_NOVO"
    assert linha["cliente_novo"] is True
    assert resultado["clientes_novos"] == 1
    assert resultado["ok"] == 1


def test_validar_planilha_feira_exige_feira(sb, tmp_path):
    _dados_validacao(sb)
    path = _planilha([{
        "Canal": "Feira", "Código Interno": "24727", "Quantidade": 1,
        "Valor Lista": 100, "Desconto": 0, "Valor Final": 100,
        "Status": "Pago", "Data": "05/10/2026", "Cliente": "Cliente Antigo",
        "Forma Pagamento": "Pix", "Feira": "",
    }], tmp_path)
    resultado = vendas.validar_planilha_vendas(path, sb)
    linha = resultado["linhas"][0]
    assert linha["status_linha"] == "ERRO"
    assert "feira" in linha["campos_erro"]
    assert any("exige uma feira" in erro for erro in linha["erros"])


def test_validar_planilha_feira_nao_e_obrigatoria_para_outro_canal(sb, tmp_path):
    _dados_validacao(sb)
    path = _planilha([{
        "Canal": "Loja", "Código Interno": "24727", "Quantidade": 1,
        "Valor Lista": 100, "Desconto": 0, "Valor Final": 100,
        "Status": "Pago", "Data": "05/10/2026", "Cliente": "Cliente Antigo",
        "Forma Pagamento": "Pix", "Feira": "Feira Centro",
    }], tmp_path)
    resultado = vendas.validar_planilha_vendas(path, sb)
    linha = resultado["linhas"][0]
    assert linha["status_linha"] == "OK"
    assert linha["feira"] == ""


def test_preparar_linhas_corrigidas_exclui_bloqueios():
    base = {
        "produto_id": 7, "canal_id": 1, "status_id": 2,
        "forma_pagamento_id": 3, "detalhe_feira_id": None,
        "quantidade": 1, "valor_lista": 100, "valor_desconto": 0,
        "valor_final": 100, "data_venda": "2026-10-05", "cliente": "A",
        "observacao": "", "duplicada": False, "erros": [],
        "status_linha": "OK",
    }
    bloqueada = {**base, "status_linha": "PRODUTO_NAO_CADASTRADO"}
    duplicada = {**base, "duplicada": True}
    prontas = vendas.preparar_linhas_corrigidas_validacao([base, bloqueada, duplicada])
    assert len(prontas) == 1
    assert prontas[0]["produto_id"] == 7
    assert prontas[0]["valor_final"] == Decimal("100")


def test_importar_linhas_validadas_cria_cliente_novo_uma_vez_e_importa_vendas(sb):
    sb.responses[("clientes", "select")] = []
    sb.responses[("clientes", "insert")] = [{"id": 90}]
    sb.responses[("parceiros", "select")] = []
    sb.responses[("vendas", "insert")] = [{"id": 50}, {"id": 51}]
    sb.responses[("estoque", "select")] = [{"quantidade_atual": 10}]
    sb.responses[("estoque", "upsert")] = [{"id": 1}]
    sb.responses[("estoque_movimentos", "insert")] = [{"id": 70}, {"id": 71}]

    linha = {
        "status_linha": "CLIENTE_NOVO", "erros": [], "motivos_nao_importar": [],
        "cliente": "Novo Cliente", "canal_id": 1, "status_id": 2,
        "forma_pagamento_id": 3, "detalhe_feira_id": None,
        "produto_id": 7, "quantidade": 1, "valor_lista": 100,
        "valor_desconto": 0, "valor_final": 100, "data_venda": "2026-10-05",
        "observacao": "",
    }
    resultado = vendas.importar_linhas_validadas([linha, linha], sb)
    assert resultado["importadas"] == 2
    assert resultado["clientes_criados"] == 1
    assert len(sb.calls_for("clientes", "insert")) == 1
    assert len(sb.calls_for("vendas", "insert")) == 2


def test_modelo_excel_contem_vendas_apoio_dropdowns_e_id(sb):
    sb.responses[("canais_venda", "select")] = [{"id": 1, "nome": "Loja"}]
    sb.responses[("formas_pagamento", "select")] = [{"id": 2, "descricao": "Pix"}]
    sb.responses[("status_venda", "select")] = [{"id": 3, "nome": "Pago"}]
    sb.responses[("detalhes_feira", "select")] = [{"id": 4, "nome_feira": "Centro"}]
    sb.responses[("clientes", "select")] = [{"id": 5, "nome": "Maria"}]
    sb.responses[("produtos", "select")] = [{"id": 7, "codigo_interno": "0024727", "descricao": "Brinco"}]
    sb.responses[("estoque", "select")] = [{"produto_id": 7, "quantidade_atual": 4}]

    arquivo = vendas.gerar_planilha_modelo_vendas(sb)
    wb = openpyxl.load_workbook(BytesIO(arquivo), data_only=False)
    assert wb.sheetnames == ["Vendas", "Dados de Apoio"]
    ws = wb["Vendas"]
    assert ws["A1"].value == "Canal"
    assert ws["B1"].value == "Código Interno"
    assert ws["F2"].value.startswith("=IF(")
    assert len(ws.data_validations.dataValidation) >= 5
    apoio = wb["Dados de Apoio"]
    assert apoio["F2"].value == "0024727"
    assert apoio["G2"].value == "Brinco"
    assert apoio["H2"].value == 4
    assert apoio["L2"].value
