"""Tela de Relatórios.

Duas abas:
  - Reports - Geral: visões de Compras, Vendas e Estoque calculadas direto do banco.
  - Reports - BI: link para o painel externo de BI (a antiga tela de relatórios).
"""
import io
import json
import os
from datetime import date, timedelta

import altair as alt
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv

from db import supabase, buscar_todos

load_dotenv()

# Compatibilidade com versões do Streamlit que usam experimental_dialog.
_dialog = getattr(st, "dialog", None) or st.experimental_dialog

_TTL = 300  # segundos de cache dos dados (botão "Atualizar dados" limpa antes)

MESES_ABREV = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

FAIXAS_IDADE = {
    "idade-verde": "0-60 dias",
    "idade-amarela": ">60-120 dias",
    "idade-laranja": ">120-180 dias",
    "idade-vermelha": ">180 dias",
}
CORES_IDADE = ["#4ade80", "#fde047", "#fb923c", "#f87171"]

# Cor padrão de cada "estilo" de relatório: todos os gráficos de barras de Compras
# usam COR_COMPRAS e todos os de Vendas usam COR_VENDAS. Para trocar o tema, basta
# alterar estes dois valores.
COR_COMPRAS = "#2563eb"  # azul
COR_VENDAS = "#16a34a"   # verde
COR_ANO_ANTERIOR = "#9ca3af"  # linha tracejada do "mesmo mês no ano anterior"
COR_ROTULO_ANTERIOR = "#4b5563"  # texto dos valores do ano anterior
COR_CLIENTE_NOVO = "#2a78d6"

# Motor dos gráficos mensais (Compras/Vendas): "echarts" (HTML interativo, carrega a
# biblioteca de um CDN) ou "altair" (versão antiga, 100% nativa do Streamlit).
GRAFICO_MENSAL_MOTOR = "echarts"

# Navegação entre Compras/Vendas/Estoque: "abas" (barra de abas com sublinhado colorido)
# ou "radio" (botões de rádio, versão antiga).
NAVEGACAO_VISOES = "abas"


# ----------------------------------------------------------------------
# Formatação e utilitários
# ----------------------------------------------------------------------
def _fmt_num(valor, casas=0) -> str:
    return (
        f"{float(valor or 0):,.{casas}f}"
        .replace(",", "X").replace(".", ",").replace("X", ".")
    )


def _fmt_moeda(valor) -> str:
    return f"R$ {_fmt_num(valor, 2)}"


def _fmt_qtd(valor) -> str:
    return f"{float(valor or 0):g}".replace(".", ",")


def _fmt_pct(valor) -> str:
    return f"{_fmt_num(valor, 1)}%"


def _fmt_curto(valor) -> str:
    valor = float(valor or 0)
    if valor >= 1000:
        return f"R$ {valor / 1000:.1f}k".replace(".", ",")
    return f"R$ {valor:.0f}"


def _fmt_data(ts) -> str:
    return "" if pd.isna(ts) else ts.strftime("%d/%m/%Y")


def _delta_pct(atual, anterior):
    """Variação percentual como texto com sinal (ex.: '+12,4%'); None sem base de comparação."""
    if not anterior:
        return None
    return f"{(atual - anterior) / anterior * 100:+.1f}%".replace(".", ",")


def _filtrar_periodo(df: pd.DataFrame, ini, fim) -> pd.DataFrame:
    """Linhas cuja coluna `data` (Timestamp) cai em [ini, fim]. Sem limites, devolve tudo."""
    mascara = pd.Series(True, index=df.index)
    if ini is not None:
        mascara &= df["data"] >= pd.Timestamp(ini)
    if fim is not None:
        mascara &= df["data"] <= pd.Timestamp(fim)
    return df[mascara]


def _num(valor) -> float:
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _parse_data(valor):
    try:
        return date.fromisoformat(str(valor or "")[:10])
    except ValueError:
        return None


def _no_periodo(dt, ini, fim) -> bool:
    if ini is None and fim is None:
        return True  # "Todo período": inclui até registros sem data
    if dt is None:
        return False
    if ini is not None and dt < ini:
        return False
    if fim is not None and dt > fim:
        return False
    return True


# ----------------------------------------------------------------------
# Carga de dados (cacheada)
# ----------------------------------------------------------------------
def _todos(tabela: str, colunas: str, ordem: str = "id"):
    return buscar_todos(lambda: supabase.table(tabela).select(colunas).order(ordem))


@st.cache_data(ttl=_TTL, show_spinner="Carregando compras...")
def _carregar_compras():
    return _todos(
        "compras",
        "id, numero_nf, data_emissao, valor_produtos, valor_desconto, valor_total, fornecedor_id",
    )


@st.cache_data(ttl=_TTL, show_spinner="Carregando itens de compra...")
def _carregar_compras_itens():
    return _todos("compras_itens", "id, produto_id, quantidade, valor_total, compras(data_emissao)")


@st.cache_data(ttl=_TTL, show_spinner=False)
def _carregar_fornecedores():
    return _todos("fornecedores", "id, razao_social, nome_fantasia")


@st.cache_data(ttl=_TTL, show_spinner=False)
def _carregar_produtos():
    return _todos("produtos", "id, codigo_interno, descricao, categoria_id")


@st.cache_data(ttl=_TTL, show_spinner=False)
def _carregar_categorias():
    return _todos("categorias_produtos", "id, categoria, banho")


@st.cache_data(ttl=_TTL, show_spinner="Carregando vendas...")
def _carregar_vendas():
    return _todos(
        "vendas",
        "id, quantidade, valor_final, comissao_valor, data_venda, cliente,"
        " produtos(id, codigo_interno, descricao, categoria_id),"
        " canais_venda(nome), status_venda(nome), formas_pagamento(descricao),"
        " detalhes_feira(nome_feira)",
    )


@st.cache_data(ttl=_TTL, show_spinner=False)
def _carregar_custos():
    """produto_id -> preço da última compra (usado como custo unitário estimado)."""
    linhas = _todos("estoque", "produto_id, estoque_preco_ultima_compra", ordem="produto_id")
    return {l["produto_id"]: _num(l.get("estoque_preco_ultima_compra")) for l in linhas}


def _rotulos_categoria() -> dict:
    return {
        c["id"]: f"{c['categoria']} | {c.get('banho') or '-'}"
        for c in _carregar_categorias()
    }


@st.cache_data(ttl=_TTL, show_spinner="Calculando estoque e idade dos produtos...")
def _carregar_estoque_visao():
    """Itens com saldo > 0, com idade média (mesmo cálculo FIFO da tela de Estoque)."""
    from telas.estoque_saldo import (
        _carregar_estoque_completo,
        _carregar_historico_compras_completo,
        _carregar_movimentos_estoque_completo,
        _calcular_idade_produto,
        _classificar_idade,
    )

    linhas = _carregar_estoque_completo()
    historico = _carregar_historico_compras_completo()
    movimentos = _carregar_movimentos_estoque_completo()

    compras_por_produto, movimentos_por_produto = {}, {}
    for c in historico:
        compras_por_produto.setdefault(c.get("produto_id"), []).append(c)
    for m in movimentos:
        movimentos_por_produto.setdefault(m.get("produto_id"), []).append(m)

    categoria_por_produto = {p["id"]: p.get("categoria_id") for p in _carregar_produtos()}
    rotulos = _rotulos_categoria()

    itens = []
    for linha in linhas:
        prod = linha.get("produtos")
        if not prod:
            continue
        saldo = _num(linha.get("quantidade_atual"))
        if saldo <= 0:
            continue
        pid = prod["id"]
        idade, _ = _calcular_idade_produto(pid, saldo, compras_por_produto, movimentos_por_produto)
        itens.append({
            "id": pid,
            "codigo": prod.get("codigo_interno") or "-",
            "produto": prod.get("descricao") or "-",
            "categoria": rotulos.get(categoria_por_produto.get(pid), "Sem categoria"),
            "saldo": saldo,
            "preco_compra": _num(linha.get("estoque_preco_ultima_compra")),
            "preco_venda": _num(linha.get("estoque_preco_venda_sugerida")),
            "idade": idade,
            "classe_idade": _classificar_idade(idade)["classe"],
        })
    return itens


# ----------------------------------------------------------------------
# Gráficos
# ----------------------------------------------------------------------
def _serie_mensal(df: pd.DataFrame, col_data: str, col_valor: str) -> pd.DataFrame:
    """Soma por mês, preenchendo meses sem movimento com zero."""
    d = df.dropna(subset=[col_data]).copy()
    if d.empty:
        return pd.DataFrame(columns=["periodo", "rotulo", "valor", "texto"])
    d["periodo"] = pd.to_datetime(d[col_data]).dt.to_period("M")
    serie = d.groupby("periodo")[col_valor].sum()
    serie = serie.reindex(pd.period_range(serie.index.min(), serie.index.max(), freq="M"), fill_value=0.0)
    out = pd.DataFrame({
        "periodo": list(serie.index),
        "rotulo": [f"{MESES_ABREV[p.month - 1]}/{str(p.year)[2:]}" for p in serie.index],
        "valor": serie.values,
    })
    out["texto"] = out["valor"].map(_fmt_moeda)
    return out


def _anexar_ano_anterior(serie: pd.DataFrame, df_base: pd.DataFrame, col_valor: str) -> pd.DataFrame:
    """Acrescenta à série mensal o valor do mesmo mês um ano antes (se houver dados).

    Meses sem movimento no ano anterior ficam vazios (a linha não despenca para zero).
    """
    if serie.empty or df_base.empty:
        return serie
    d = df_base.dropna(subset=["data"])
    por_mes = d.groupby(d["data"].dt.to_period("M"))[col_valor].sum()
    anterior = [float(por_mes.get(p - 12, 0.0)) or None for p in serie["periodo"]]
    if not any(anterior):
        return serie
    serie = serie.copy()
    serie["anterior"] = pd.Series(anterior, index=serie.index, dtype="float64")
    serie["texto_anterior"] = [_fmt_moeda(a) if a else "-" for a in anterior]
    serie["rotulo_anterior"] = [_fmt_curto(a) if a else "" for a in anterior]
    serie["variacao"] = [_delta_pct(v, a) or "-" for v, a in zip(serie["valor"], anterior)]
    return serie


def _legenda_mensal(nome_serie: str, cor: str):
    """Legenda do gráfico mensal com comparação: barras = período, tracejado = ano anterior."""
    st.markdown(
        '<div style="display:flex;flex-wrap:wrap;gap:20px;align-items:center;'
        'font-size:0.85rem;margin:0 0 0.35rem 0">'
        f'<span><span style="display:inline-block;width:12px;height:12px;border-radius:3px;'
        f'background:{cor};margin-right:6px;vertical-align:-1px"></span>{nome_serie}</span>'
        f'<span><span style="display:inline-block;width:22px;border-top:2px dashed {COR_ANO_ANTERIOR};'
        f'margin-right:6px;vertical-align:3px"></span>Mesmo mês do ano anterior</span>'
        "</div>",
        unsafe_allow_html=True,
    )


def _chart_mensal(df: pd.DataFrame, titulo_valor: str, cor: str) -> alt.Chart:
    """Barras por mês com o valor exato (R$) em cima de cada barra, a 45°.

    O texto usa a própria cor da série, em negrito, com um contorno claro por
    trás (legível nos temas claro e escuro e sobre as barras vizinhas). Se a
    série tiver a coluna `anterior`, desenha também a linha tracejada do mesmo
    mês no ano anterior, com o valor (compacto) logo abaixo de cada ponto, e o
    tooltip das barras passa a mostrar o ano anterior e a variação.
    """
    df = df.drop(columns=["periodo"], errors="ignore")
    tem_anterior = "anterior" in df.columns
    maior = float(df["valor"].max())
    if tem_anterior:
        maior = max(maior, float(df["anterior"].max()))
    escala = alt.Scale(domain=[0, max(maior * 1.4, 1.0)])  # folga no alto para os labels

    base = alt.Chart(df).encode(
        x=alt.X("rotulo:N", sort=df["rotulo"].tolist(), title=None, axis=alt.Axis(labelAngle=-45)),
        y=alt.Y("valor:Q", title=None, axis=alt.Axis(format="~s"), scale=escala),
    )
    tooltip = [alt.Tooltip("rotulo:N", title="Mês"), alt.Tooltip("texto:N", title=titulo_valor)]
    if tem_anterior:
        tooltip += [
            alt.Tooltip("texto_anterior:N", title="Ano anterior"),
            alt.Tooltip("variacao:N", title="Variação"),
        ]
    camadas = [base.mark_bar(color=cor).encode(tooltip=tooltip)]

    if tem_anterior:
        y_ant = alt.Y("anterior:Q", title=None, axis=alt.Axis(format="~s"), scale=escala)
        camadas.append(
            base.mark_line(
                color=COR_ANO_ANTERIOR, strokeDash=[5, 4], strokeWidth=2,
                point=alt.OverlayMarkDef(color=COR_ANO_ANTERIOR, size=45),
            ).encode(y=y_ant)
        )
        # Valor do ano anterior (compacto) junto a cada ponto. Ponto ACIMA da barra: texto à
        # ESQUERDA do ponto (o label inclinado da barra sobe para a direita, então a esquerda
        # fica livre). Ponto ABAIXO do topo da barra: texto embaixo do ponto, dentro da barra.
        posicoes = (
            ("isValid(datum.anterior) && datum.anterior >= datum.valor",
             dict(align="right", baseline="middle", dx=-9)),
            ("isValid(datum.anterior) && datum.anterior < datum.valor",
             dict(align="center", baseline="top", dy=9)),
        )
        for condicao, posicao in posicoes:
            for kwargs in (
                dict(color="white", stroke="white", strokeWidth=3, strokeJoin="round"),  # contorno
                dict(color=COR_ROTULO_ANTERIOR),
            ):
                camadas.append(
                    base.mark_text(fontSize=12, fontWeight="bold", **posicao, **kwargs)
                    .encode(y=y_ant, text="rotulo_anterior:N")
                    .transform_filter(condicao)
                )

    estilo_texto = dict(
        angle=315, align="left", baseline="middle", dx=6, fontSize=14, fontWeight="bold",
    )
    # Contorno claro atrás do texto: a 45° o label de uma barra baixa pode passar por
    # cima da barra vizinha (mais alta); o contorno mantém o texto legível sobre ela.
    camadas.append(
        base.mark_text(
            color="white", stroke="white", strokeWidth=3, strokeJoin="round", **estilo_texto
        ).encode(text="texto:N")
    )
    camadas.append(base.mark_text(color=cor, **estilo_texto).encode(text="texto:N"))
    # padding à direita: o label da última barra se estende para a direita.
    return alt.layer(*camadas).properties(
        height=340, padding={"left": 5, "top": 10, "right": 80, "bottom": 5}
    )


_ECHARTS_URL = "https://cdn.jsdelivr.net/npm/echarts@5.5.0/dist/echarts.min.js"

_HTML_GRAFICO_MENSAL = """<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
  html, body { margin: 0; padding: 0; background: transparent;
    font-family: "Source Sans Pro", -apple-system, "Segoe UI", Roboto, sans-serif; }
  #g { width: 100%; height: __ALTURA__px; }
  .erro { padding: 24px; font-size: 14px; color: #6b7280; }
</style>
<script src="__ECHARTS_URL__"></script>
</head><body>
<div id="g"></div>
<script>
const D = __DADOS__;
(function () {
  const el = document.getElementById('g');
  if (typeof echarts === 'undefined') {
    el.innerHTML = '<div class="erro">Não foi possível carregar a biblioteca do gráfico ' +
      '(ECharts). Verifique a conexão com a internet.</div>';
    return;
  }
  const escuro = (D.tema || (matchMedia('(prefers-color-scheme: dark)').matches ? 'dark' : 'light')) === 'dark';
  const C = escuro
    ? { texto: '#e5e7eb', sub: '#9ca3af', linha: 'rgba(255,255,255,0.12)', sombra: 'rgba(255,255,255,0.06)',
        ant: '#9ca3af', tipFundo: '#1f2937', tipBorda: '#374151', ponto: '#0e1117' }
    : { texto: '#1f2937', sub: '#6b7280', linha: 'rgba(0,0,0,0.10)', sombra: 'rgba(0,0,0,0.04)',
        ant: '#9ca3af', tipFundo: '#ffffff', tipBorda: '#e5e7eb', ponto: '#ffffff' };

  const nf = (n, d) => n.toLocaleString('pt-BR', { minimumFractionDigits: 0, maximumFractionDigits: d });
  const compacto = (v) => {
    const a = Math.abs(v);
    if (a >= 1e6) return nf(v / 1e6, 2) + ' mi';
    if (a >= 1e3) return nf(v / 1e3, 1) + ' mil';
    return nf(v, 0);
  };

  const n = D.rotulos.length;
  const temAnt = Array.isArray(D.anterior);
  const zoom = n > 18;
  const NOME_ANT = 'Mesmo mês do ano anterior';

  const barras = D.valores.map((v, i) => ({
    value: v,
    itemStyle: D.parcial[i] ? { opacity: 0.55 } : {}
  }));

  const series = [{
    name: D.nome, type: 'bar', data: barras, barMaxWidth: 48, z: 2,
    itemStyle: {
      borderRadius: [6, 6, 0, 0],
      color: new echarts.graphic.LinearGradient(0, 0, 0, 1,
        [{ offset: 0, color: D.cor }, { offset: 1, color: D.cor + 'b3' }])
    },
    label: {
      show: true, position: 'top', distance: 4, color: C.texto, fontSize: 12, fontWeight: 600,
      formatter: (p) => (p.value ? compacto(p.value) : '')
    },
    labelLayout: { hideOverlap: true },
    emphasis: { itemStyle: { shadowBlur: 12, shadowColor: D.cor + '66' } }
  }];

  if (temAnt) {
    series.push({
      name: NOME_ANT, type: 'line', data: D.anterior, z: 3,
      symbol: 'circle', symbolSize: 7, connectNulls: false,
      lineStyle: { color: C.ant, width: 2, type: 'dashed' },
      itemStyle: { color: C.ant, borderColor: C.ponto, borderWidth: 2 }
    });
  }

  const linha = (cor, nome, texto) =>
    '<div style="display:flex;justify-content:space-between;gap:18px;margin-top:3px">' +
    '<span><span style="display:inline-block;width:9px;height:9px;border-radius:50%;background:' + cor +
    ';margin-right:6px"></span>' + nome + '</span><b>' + texto + '</b></div>';

  const legenda = [{ name: D.nome, itemStyle: { color: D.cor } }];
  if (temAnt) legenda.push({ name: NOME_ANT, itemStyle: { color: C.ant }, lineStyle: { color: C.ant, type: 'dashed' } });

  const chart = echarts.init(el, null, { renderer: 'canvas' });
  chart.setOption({
    animationDuration: 600,
    textStyle: { fontFamily: '"Source Sans Pro", -apple-system, "Segoe UI", Roboto, sans-serif' },
    legend: { top: 0, left: 0, selectedMode: false, itemWidth: 14, itemHeight: 10, itemGap: 22,
              textStyle: { color: C.texto, fontSize: 13 }, data: legenda },
    grid: { left: 8, right: 14, top: 44, bottom: zoom ? 58 : 8, containLabel: true },
    tooltip: {
      trigger: 'axis', axisPointer: { type: 'shadow', shadowStyle: { color: C.sombra } },
      backgroundColor: C.tipFundo, borderColor: C.tipBorda, padding: [10, 12],
      textStyle: { color: C.texto, fontSize: 13 },
      formatter: (ps) => {
        const i = ps[0].dataIndex;
        let h = '<div style="font-weight:700;margin-bottom:5px">' + D.rotulos[i] +
          (D.parcial[i] ? ' <span style="font-weight:400;color:' + C.sub + '">(em andamento)</span>' : '') + '</div>';
        h += linha(D.cor, D.nome, D.textos[i]);
        if (temAnt) {
          h += linha(C.ant, 'Ano anterior', D.textos_ant[i]);
          const v = D.variacoes[i];
          if (v && v !== '-') {
            h += '<div style="margin-top:7px;color:' + C.sub + '">Variação: <b style="color:' + C.texto + '">' +
              (v.charAt(0) === '-' ? '▼ ' : '▲ ') + v.replace(/^[+-]/, '') + '</b></div>';
          }
        }
        return h;
      }
    },
    xAxis: { type: 'category', data: D.rotulos, axisTick: { show: false },
             axisLine: { lineStyle: { color: C.linha } },
             axisLabel: { color: C.sub, fontSize: 12, hideOverlap: true } },
    yAxis: { type: 'value', name: 'R$', nameTextStyle: { color: C.sub, align: 'left' },
             splitLine: { lineStyle: { color: C.linha, type: 'dashed' } },
             axisLabel: { color: C.sub, formatter: compacto } },
    dataZoom: zoom ? [
      { type: 'inside', startValue: n - 18, endValue: n - 1,
        zoomOnMouseWheel: false, moveOnMouseWheel: false, moveOnMouseMove: true },
      { type: 'slider', startValue: n - 18, endValue: n - 1, height: 18, bottom: 6,
        borderColor: 'transparent', backgroundColor: C.sombra, fillerColor: D.cor + '33',
        handleSize: '80%', brushSelect: false, textStyle: { color: C.sub } }
    ] : [],
    series: series
  });

  new ResizeObserver(() => chart.resize()).observe(el);
})();
</script>
</body></html>"""


def _tema_streamlit():
    """'light' / 'dark' conforme o tema atual do Streamlit (None = deixar o
    navegador decidir). O gráfico roda num iframe e não herda o tema sozinho."""
    try:
        tipo = st.context.theme.type  # Streamlit >= 1.46
        if tipo in ("light", "dark"):
            return tipo
    except Exception:
        pass
    try:
        base = st.get_option("theme.base")
        if base in ("light", "dark"):
            return base
    except Exception:
        pass
    return None


def _html_grafico_mensal(df: pd.DataFrame, titulo_valor: str, cor: str, tema=None) -> tuple[str, int]:
    """Monta o HTML (ECharts) do gráfico mensal e devolve (html, altura_px).

    Barras = valor do mês, com o valor compacto em cima (R$ 12,3 mil); linha
    tracejada = mesmo mês do ano anterior (valor exato e variação no
    tooltip). O mês corrente fica com a barra mais clara ("em andamento").
    Com mais de 18 meses aparece uma barra de rolagem para navegar no tempo.
    """
    mes_corrente = pd.Period(date.today(), freq="M")
    periodos = df["periodo"].tolist() if "periodo" in df.columns else [None] * len(df)

    def _limpo(v):
        return None if v is None or pd.isna(v) else float(v)

    dados = {
        "rotulos": df["rotulo"].tolist(),
        "valores": [float(v) for v in df["valor"]],
        "textos": df["texto"].tolist(),
        "parcial": [p == mes_corrente for p in periodos],
        "nome": titulo_valor,
        "cor": cor,
        "tema": tema,
        "anterior": None,
    }
    if "anterior" in df.columns:
        dados["anterior"] = [_limpo(v) for v in df["anterior"]]
        dados["textos_ant"] = df["texto_anterior"].tolist()
        dados["variacoes"] = df["variacao"].tolist()

    altura = 360 + (36 if len(df) > 18 else 0)
    # "</" dentro do JSON fecharia a tag <script> antes da hora.
    payload = json.dumps(dados, ensure_ascii=False).replace("</", "<\\/")
    html = (
        _HTML_GRAFICO_MENSAL
        .replace("__ECHARTS_URL__", _ECHARTS_URL)
        .replace("__ALTURA__", str(altura))
        .replace("__DADOS__", payload)
    )
    return html, altura


def _grafico_mensal(df: pd.DataFrame, titulo_valor: str, cor: str):
    if df.empty:
        st.caption("Sem dados com data no período.")
        return

    if GRAFICO_MENSAL_MOTOR == "echarts":
        try:
            html, altura = _html_grafico_mensal(df, titulo_valor, cor, _tema_streamlit())
            components.html(html, height=altura + 4)
            return
        except Exception as err:
            st.caption(f"Gráfico interativo indisponível ({err}); exibindo versão simples.")

    # Versão Altair (fallback ou GRAFICO_MENSAL_MOTOR = "altair").
    if "anterior" in df.columns:
        _legenda_mensal(titulo_valor, cor)
    st.altair_chart(_chart_mensal(df, titulo_valor, cor))


def _top_agrupado(df: pd.DataFrame, col_grupo: str, col_valor: str, n: int = 10) -> pd.DataFrame:
    g = df.groupby(col_grupo)[col_valor].sum().sort_values(ascending=False).head(n)
    out = pd.DataFrame({"rotulo": g.index.astype(str), "valor": g.values})
    out["texto"] = out["valor"].map(_fmt_moeda)
    return out


def _grafico_horizontal(df: pd.DataFrame, titulo_valor: str, cor: str | None = None):
    if df.empty:
        st.caption("Sem dados no período.")
        return
    chart = alt.Chart(df).mark_bar(**({"color": cor} if cor else {})).encode(
        y=alt.Y("rotulo:N", sort="-x", title=None, axis=alt.Axis(labelLimit=240)),
        x=alt.X("valor:Q", title=None, axis=alt.Axis(format="~s")),
        tooltip=[alt.Tooltip("rotulo:N", title="Item"), alt.Tooltip("texto:N", title=titulo_valor)],
    ).properties(height=max(140, 30 * len(df) + 20))
    st.altair_chart(chart)


def _kpi(coluna, rotulo: str, valor: str, ajuda: str | None = None, detalhe: str | None = None,
         cor_delta: str = "off"):
    with coluna:
        with st.container(border=True):
            st.metric(rotulo, valor, delta=detalhe, delta_color=cor_delta, help=ajuda)


# ----------------------------------------------------------------------
# Visão: Compras
# ----------------------------------------------------------------------
def _visao_compras(ini, fim):
    try:
        fornecedores = {
            f["id"]: f.get("nome_fantasia") or f.get("razao_social") or "-"
            for f in _carregar_fornecedores()
        }
        produtos = {p["id"]: p for p in _carregar_produtos()}

        linhas = []
        for c in _carregar_compras():
            linhas.append({
                "data": _parse_data(c.get("data_emissao")),
                "fornecedor": fornecedores.get(c.get("fornecedor_id"), "Sem fornecedor"),
                "valor_produtos": _num(c.get("valor_produtos")),
                "desconto": _num(c.get("valor_desconto")),
                "valor_total": _num(c.get("valor_total")),
            })
        todas = pd.DataFrame(linhas, columns=["data", "fornecedor", "valor_produtos", "desconto", "valor_total"])
        todas["data"] = pd.to_datetime(todas["data"])
        df = _filtrar_periodo(todas, ini, fim)

        itens = []
        for it in _carregar_compras_itens():
            dt = _parse_data((it.get("compras") or {}).get("data_emissao"))
            if not _no_periodo(dt, ini, fim):
                continue
            prod = produtos.get(it.get("produto_id"), {})
            itens.append({
                "codigo": prod.get("codigo_interno") or "-",
                "produto": prod.get("descricao") or "-",
                "quantidade": _num(it.get("quantidade")),
                "valor_total": _num(it.get("valor_total")),
            })
        df_it = pd.DataFrame(itens, columns=["codigo", "produto", "quantidade", "valor_total"])
    except Exception as e:
        st.error(f"Erro ao carregar dados de compras: {e}")
        return

    if df.empty:
        st.info("Nenhuma compra no período selecionado.")
        return

    c1, c2, c3 = st.columns(3)
    _kpi(c1, "Notas fiscais", _fmt_num(len(df)))
    _kpi(c2, "Valor comprado", _fmt_moeda(df["valor_total"].sum()))
    _kpi(c3, "Peças compradas", _fmt_qtd(df_it["quantidade"].sum()))

    st.markdown("---")
    st.subheader("Valor comprado por mês")
    serie = _anexar_ano_anterior(_serie_mensal(df, "data", "valor_total"), todas, "valor_total")
    _grafico_mensal(serie, "Valor comprado", COR_COMPRAS)

    st.subheader("Produtos mais comprados")
    if df_it.empty:
        st.caption("Sem itens de compra no período.")
        return
    g = (
        df_it.groupby(["codigo", "produto"], as_index=False)
        .agg(pecas=("quantidade", "sum"), valor=("valor_total", "sum"))
        .sort_values("pecas", ascending=False)
        .head(10)
    )
    st.dataframe(
        pd.DataFrame({
            "Cód. Produto": g["codigo"],
            "Produto": g["produto"],
            "Peças": g["pecas"].map(_fmt_qtd),
            "Valor": g["valor"].map(_fmt_moeda),
            "Preço médio": [_fmt_moeda(v / p) if p else "-" for v, p in zip(g["valor"], g["pecas"])],
        }),
        hide_index=True,
        use_container_width=True,
    )


# ----------------------------------------------------------------------
# Visão: Vendas
# ----------------------------------------------------------------------
_COLUNAS_VENDAS = [
    "id", "data", "quantidade", "valor_final", "comissao", "canal", "forma", "status",
    "codigo", "produto", "categoria", "custo", "tem_custo", "cliente", "cliente_key", "feira",
]


@st.cache_data(ttl=_TTL, show_spinner="Carregando vendas...")
def _df_vendas_base() -> pd.DataFrame:
    """Todas as vendas em um DataFrame (uma linha por venda), sem filtros."""
    categorias = _rotulos_categoria()
    custos = _carregar_custos()

    linhas = []
    for v in _carregar_vendas():
        prod = v.get("produtos") or {}
        qtd = _num(v.get("quantidade"))
        custo_unit = custos.get(prod.get("id"), 0.0)
        cliente = str(v.get("cliente") or "").strip()
        linhas.append({
            "id": v.get("id"),
            "data": _parse_data(v.get("data_venda")),
            "quantidade": qtd,
            "valor_final": _num(v.get("valor_final")),
            "comissao": _num(v.get("comissao_valor")),
            "canal": (v.get("canais_venda") or {}).get("nome") or "Sem canal",
            "forma": (v.get("formas_pagamento") or {}).get("descricao") or "Não informado",
            "status": (v.get("status_venda") or {}).get("nome") or "Sem status",
            "codigo": prod.get("codigo_interno") or "-",
            "produto": prod.get("descricao") or "-",
            "categoria": categorias.get(prod.get("categoria_id"), "Sem categoria"),
            "custo": custo_unit * qtd,
            "tem_custo": custo_unit > 0,
            "cliente": cliente,
            "cliente_key": cliente.lower(),
            "feira": (v.get("detalhes_feira") or {}).get("nome_feira") or "",
        })
    df = pd.DataFrame(linhas, columns=_COLUNAS_VENDAS)
    df["data"] = pd.to_datetime(df["data"])
    return df


def _filtrar_vendas(df: pd.DataFrame, ini, fim, canal="Todos", categoria="Todas", status="Todos") -> pd.DataFrame:
    df = _filtrar_periodo(df, ini, fim)
    mascara = pd.Series(True, index=df.index)
    if canal != "Todos":
        mascara &= df["canal"] == canal
    if categoria != "Todas":
        mascara &= df["categoria"] == categoria
    if status != "Todos":
        mascara &= df["status"] == status
    return df[mascara]


def _resumo_vendas(df: pd.DataFrame) -> dict:
    n = len(df)
    fat = float(df["valor_final"].sum())
    com_custo = df[df["tem_custo"]]
    fat_cc = float(com_custo["valor_final"].sum())
    margem = fat_cc - float(com_custo["custo"].sum())
    return {
        "n": n,
        "fat": fat,
        "ticket": fat / n if n else 0.0,
        "pecas": float(df["quantidade"].sum()),
        "comissoes": float(df["comissao"].sum()),
        "margem": margem,
        "margem_pct": margem / fat_cc * 100 if fat_cc else 0.0,
        "cobertura": fat_cc / fat * 100 if fat else 0.0,
    }


def _gerar_excel_vendas(df: pd.DataFrame) -> bytes:
    """Planilha com as vendas filtradas (uma linha por venda) e linha de total."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"

    cabecalhos = [
        "Data", "Cliente", "Cód. Produto", "Produto", "Categoria", "Canal", "Feira",
        "Forma de Pagamento", "Status", "Qtde", "Valor Final", "Comissão",
    ]
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for col, titulo in enumerate(cabecalhos, 1):
        c = ws.cell(1, col, titulo)
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    linha = 2
    for r in df.sort_values(["data", "id"], ascending=False).itertuples():
        valores = [
            None if pd.isna(r.data) else r.data.date(),
            r.cliente, r.codigo, r.produto, r.categoria, r.canal, r.feira,
            r.forma, r.status, float(r.quantidade), float(r.valor_final), float(r.comissao),
        ]
        for col, valor in enumerate(valores, 1):
            ws.cell(linha, col, valor)
        ws.cell(linha, 1).number_format = "DD/MM/YYYY"
        ws.cell(linha, 3).number_format = "@"
        ws.cell(linha, 10).number_format = "0.##"
        ws.cell(linha, 11).number_format = "R$ #,##0.00"
        ws.cell(linha, 12).number_format = "R$ #,##0.00"
        linha += 1

    ultima = linha - 1
    ws.cell(linha, 9, "TOTAL").font = Font(bold=True)
    ws.cell(linha, 9).alignment = Alignment(horizontal="right")
    for col in range(1, len(cabecalhos) + 1):
        ws.cell(linha, col).border = Border(top=Side(style="thin"))
    for col, formato in ((10, "0.##"), (11, "R$ #,##0.00"), (12, "R$ #,##0.00")):
        letra = get_column_letter(col)
        c = ws.cell(linha, col, f"=SUBTOTAL(109,{letra}2:{letra}{ultima})")
        c.font = Font(bold=True)
        c.number_format = formato

    for i, largura in enumerate([12, 24, 14, 36, 22, 14, 22, 20, 14, 8, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecalhos))}{ultima}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@_dialog("📥 Exportar vendas", width="small")
def _dialog_exportar_vendas(df_exportar: pd.DataFrame, descricao_filtros: str):
    """Download em Excel das vendas com os filtros atuais (gerado só ao abrir)."""
    st.caption(f"Filtros aplicados: **{descricao_filtros}**")

    if df_exportar.empty:
        st.info("Nenhuma venda para exportar com os filtros atuais.")
        return

    try:
        with st.spinner("Gerando planilha..."):
            dados = _gerar_excel_vendas(df_exportar)
    except Exception as e:
        st.error(f"Não foi possível gerar a planilha: {e}")
        return

    st.success(f"{len(df_exportar)} venda(s) prontas para exportar.")
    st.download_button(
        "⬇️ Baixar Excel",
        data=dados,
        file_name=f"vendas_{date.today():%Y%m%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="rep_v_download_excel",
    )


def _bloco_kpis_vendas(atual: dict, anterior: dict | None):
    """Seis KPIs; com `anterior`, mostra a variação contra o período anterior."""
    def d_pct(chave):
        return _delta_pct(atual[chave], anterior[chave]) if anterior else None

    delta_margem = None
    if anterior and anterior["margem_pct"]:
        delta_margem = f"{atual['margem_pct'] - anterior['margem_pct']:+.1f}".replace(".", ",") + " p.p."

    c1, c2, c3 = st.columns(3)
    _kpi(c1, "Faturamento", _fmt_moeda(atual["fat"]), detalhe=d_pct("fat"), cor_delta="normal")
    _kpi(c2, "Vendas", _fmt_num(atual["n"]), detalhe=d_pct("n"), cor_delta="normal")
    _kpi(c3, "Ticket médio", _fmt_moeda(atual["ticket"]), detalhe=d_pct("ticket"), cor_delta="normal")
    c4, c5, c6 = st.columns(3)
    _kpi(c4, "Peças vendidas", _fmt_qtd(atual["pecas"]), detalhe=d_pct("pecas"), cor_delta="normal")
    _kpi(
        c5, "Margem bruta estimada", _fmt_moeda(atual["margem"]),
        ajuda=(
            "Faturamento menos o custo estimado (quantidade vendida × preço da última "
            "compra do produto). Só considera vendas de produtos com preço de compra "
            f"conhecido ({_fmt_pct(atual['cobertura'])} do faturamento do período). "
            "A variação é da margem percentual, em pontos percentuais (p.p.)."
        ),
        detalhe=delta_margem or f"{_fmt_pct(atual['margem_pct'])} sobre o faturamento com custo",
        cor_delta="normal" if delta_margem else "off",
    )
    _kpi(c6, "Comissões", _fmt_moeda(atual["comissoes"]), detalhe=d_pct("comissoes"))


def _bloco_clientes(df: pd.DataFrame, base: pd.DataFrame, ini):
    """Novos x recorrentes (1ª compra dentro do período = novo) e top clientes."""
    com = df[df["cliente_key"] != ""].copy()
    if com.empty:
        st.caption("Nenhuma venda com cliente informado no período.")
        return

    if ini is not None:
        primeira = base[base["cliente_key"] != ""].groupby("cliente_key")["data"].min()
        com["novo"] = com["cliente_key"].map(primeira) >= pd.Timestamp(ini)
        por_cliente = com.groupby("cliente_key").agg(novo=("novo", "first"))
        n_novos = int(por_cliente["novo"].sum())
        n_rec = len(por_cliente) - n_novos

        dados = pd.DataFrame({"grupo": ["Novos", "Recorrentes"], "clientes": [n_novos, n_rec]})
        dados["pct"] = dados["clientes"] / max(len(por_cliente), 1) * 100
        dados["texto"] = [f"{c} ({_fmt_pct(p)})" for c, p in zip(dados["clientes"], dados["pct"])]
        st.altair_chart(
            alt.Chart(dados).mark_bar().encode(
                x=alt.X("clientes:Q", stack="normalize", axis=None),
                color=alt.Color(
                    "grupo:N",
                    scale=alt.Scale(domain=["Novos", "Recorrentes"], range=[COR_CLIENTE_NOVO, COR_VENDAS]),
                    legend=alt.Legend(orient="bottom", title=None),
                ),
                tooltip=[alt.Tooltip("grupo:N", title="Grupo"), alt.Tooltip("texto:N", title="Clientes")],
            ).properties(height=50)
        )

        def ticket(novo: bool) -> str:
            sub = com[com["novo"] == novo]
            return _fmt_moeda(sub["valor_final"].sum() / len(sub)) if len(sub) else "-"

        st.caption(f"Ticket médio: novos **{ticket(True)}** · recorrentes **{ticket(False)}**")
    else:
        st.caption("Escolha um período (que não seja “Todo período”) para ver novos x recorrentes.")

    top = (
        com.groupby("cliente_key")
        .agg(cliente=("cliente", "first"), compras=("id", "count"), fat=("valor_final", "sum"))
        .sort_values("fat", ascending=False)
        .head(5)
    )
    st.dataframe(
        pd.DataFrame({
            "Cliente": top["cliente"],
            "Compras": top["compras"].map(_fmt_num),
            "Faturamento": top["fat"].map(_fmt_moeda),
        }),
        hide_index=True,
        use_container_width=True,
    )


def _bloco_feiras(df: pd.DataFrame):
    com = df[df["feira"] != ""]
    if com.empty:
        st.caption("Nenhuma venda associada a feira/evento no período.")
        return
    g = (
        com.groupby("feira")
        .agg(vendas=("id", "count"), fat=("valor_final", "sum"))
        .sort_values("fat", ascending=False)
        .head(5)
    )
    st.dataframe(
        pd.DataFrame({
            "Evento": g.index,
            "Vendas": g["vendas"].map(_fmt_num),
            "Faturamento": g["fat"].map(_fmt_moeda),
        }),
        hide_index=True,
        use_container_width=True,
    )


def _visao_vendas(ini, fim):
    try:
        base = _df_vendas_base()
    except Exception as e:
        st.error(f"Erro ao carregar dados de vendas: {e}")
        return
    if base.empty:
        st.info("Nenhuma venda cadastrada.")
        return

    # --- Filtros da visão ---------------------------------------------------
    col_canal, col_cat, col_status, col_comp, col_exp = st.columns(
        [1.3, 1.8, 1.3, 1.8, 1.1], vertical_alignment="bottom"
    )
    with col_canal:
        canal = st.selectbox("Canal", ["Todos"] + sorted(base["canal"].unique()), key="rep_v_canal")
    with col_cat:
        categoria = st.selectbox("Categoria", ["Todas"] + sorted(base["categoria"].unique()), key="rep_v_categoria")
    with col_status:
        status = st.selectbox("Status", ["Todos"] + sorted(base["status"].unique()), key="rep_v_status")
    with col_comp:
        comparar = st.toggle(
            "Comparar com período anterior",
            value=True,
            key="rep_v_comparar",
            disabled=ini is None,
            help="Compara com o período imediatamente anterior, de mesma duração. "
                 "Indisponível em “Todo período”.",
        )
    with col_exp:
        clicou_exportar = st.button("📥 Exportar", use_container_width=True, key="rep_v_exportar")
    comparar = comparar and ini is not None

    df = _filtrar_vendas(base, ini, fim, canal, categoria, status)
    sem_periodo = _filtrar_vendas(base, None, None, canal, categoria, status)

    if clicou_exportar:
        filtros = ["Todo período" if ini is None else f"{ini:%d/%m/%Y} a {fim:%d/%m/%Y}"]
        if canal != "Todos":
            filtros.append(f"Canal: {canal}")
        if categoria != "Todas":
            filtros.append(f"Categoria: {categoria}")
        if status != "Todos":
            filtros.append(f"Status: {status}")
        _dialog_exportar_vendas(df, " · ".join(filtros))

    if df.empty:
        st.info("Nenhuma venda para os filtros selecionados.")
        return

    # --- Período anterior (mesma duração, logo antes do atual) --------------
    anterior = None
    if comparar:
        dias = (fim - ini).days + 1
        fim_ant = ini - timedelta(days=1)
        ini_ant = fim_ant - timedelta(days=dias - 1)
        df_ant = _filtrar_vendas(base, ini_ant, fim_ant, canal, categoria, status)
        if df_ant.empty:
            st.caption(f"Sem vendas em {ini_ant:%d/%m/%Y} a {fim_ant:%d/%m/%Y} para comparar.")
        else:
            anterior = _resumo_vendas(df_ant)
            st.caption(
                f"Variações em relação a {ini_ant:%d/%m/%Y} a {fim_ant:%d/%m/%Y} (período anterior)."
            )

    # --- KPIs ---------------------------------------------------------------
    _bloco_kpis_vendas(_resumo_vendas(df), anterior)

    # --- Faturamento por mês (com linha do ano anterior) --------------------
    st.markdown("---")
    st.subheader("Faturamento por mês")
    serie = _anexar_ano_anterior(_serie_mensal(df, "data", "valor_final"), sem_periodo, "valor_final")
    _grafico_mensal(serie, "Faturamento", COR_VENDAS)

    # --- Quebras ------------------------------------------------------------
    col_c, col_d, col_e = st.columns(3)
    with col_c:
        st.subheader("Por canal de venda")
        _grafico_horizontal(_top_agrupado(df, "canal", "valor_final", 8), "Faturamento", COR_VENDAS)
    with col_d:
        st.subheader("Por categoria")
        _grafico_horizontal(_top_agrupado(df, "categoria", "valor_final", 10), "Faturamento", COR_VENDAS)
    with col_e:
        st.subheader("Por forma de pagamento")
        _grafico_horizontal(_top_agrupado(df, "forma", "valor_final", 8), "Faturamento", COR_VENDAS)

    # --- Clientes e feiras --------------------------------------------------
    col_cli, col_fei = st.columns(2)
    with col_cli:
        st.subheader("Clientes")
        _bloco_clientes(df, base, ini)
    with col_fei:
        st.subheader("Feiras e eventos")
        _bloco_feiras(df)

    # --- Últimas vendas -----------------------------------------------------
    st.subheader("Últimas vendas")
    st.caption("As 10 mais recentes do período e filtros selecionados.")
    ultimas = df.sort_values(["data", "id"], ascending=False).head(10)
    st.dataframe(
        pd.DataFrame({
            "Data": ultimas["data"].map(_fmt_data),
            "Cliente": ultimas["cliente"].replace("", "-"),
            "Produto": ultimas["produto"],
            "Canal": ultimas["canal"],
            "Valor": ultimas["valor_final"].map(_fmt_moeda),
            "Status": ultimas["status"],
        }),
        hide_index=True,
        use_container_width=True,
    )


# ----------------------------------------------------------------------
# Visão: Estoque
# ----------------------------------------------------------------------
def _visao_estoque():
    st.caption("O estoque é uma foto de hoje (somente produtos com saldo maior que zero); o filtro de período não se aplica.")
    try:
        itens = _carregar_estoque_visao()
    except ImportError as e:
        st.error(
            "Não consegui importar os cálculos da tela de Estoque (`telas/estoque_saldo.py`): "
            f"{e}. Ajuste o caminho do import em `_carregar_estoque_visao`."
        )
        return
    except Exception as e:
        st.error(f"Erro ao carregar estoque: {e}")
        return

    if not itens:
        st.info("Nenhum produto com estoque maior que zero.")
        return

    pecas = sum(i["saldo"] for i in itens)
    custo_total = sum(i["saldo"] * i["preco_compra"] for i in itens)
    venda_total = sum(i["saldo"] * i["preco_venda"] for i in itens)
    margem = venda_total - custo_total
    margem_pct = margem / venda_total * 100 if venda_total else 0
    parado = sum(i["saldo"] * i["preco_compra"] for i in itens if i["classe_idade"] == "idade-vermelha")
    parado_pct = parado / custo_total * 100 if custo_total else 0

    c1, c2, c3 = st.columns(3)
    _kpi(c1, "Peças em estoque", _fmt_qtd(pecas))
    _kpi(c2, "Produtos distintos", _fmt_num(len(itens)))
    _kpi(c3, "Valor a custo (última compra)", _fmt_moeda(custo_total))
    c4, c5, c6 = st.columns(3)
    _kpi(c4, "Valor a venda sugerida", _fmt_moeda(venda_total))
    _kpi(c5, "Margem potencial", _fmt_moeda(margem), detalhe=f"{_fmt_pct(margem_pct)} sobre a venda sugerida")
    _kpi(
        c6, "Capital parado (>180 dias)", _fmt_moeda(parado),
        ajuda="Valor a custo dos produtos cuja idade média do estoque passa de 180 dias.",
        detalhe=f"{_fmt_pct(parado_pct)} do valor a custo",
    )

    sem_custo = sum(1 for i in itens if i["preco_compra"] <= 0)
    if sem_custo:
        st.caption(f"⚠️ {sem_custo} produto(s) em estoque estão sem preço de última compra e entram com custo zero.")

    st.markdown("---")
    col_a, col_b = st.columns(2)

    with col_a:
        st.subheader("Valor a custo por idade")
        linhas = []
        for classe, rotulo in FAIXAS_IDADE.items():
            grupo = [i for i in itens if i["classe_idade"] == classe]
            valor = sum(i["saldo"] * i["preco_compra"] for i in grupo)
            linhas.append({
                "faixa": rotulo,
                "valor": valor,
                "texto": _fmt_moeda(valor),
                "detalhe": f"{len(grupo)} produtos · {_fmt_qtd(sum(i['saldo'] for i in grupo))} peças",
            })
        df_idade = pd.DataFrame(linhas)
        chart = alt.Chart(df_idade).mark_bar().encode(
            x=alt.X("faixa:N", sort=list(FAIXAS_IDADE.values()), title=None, axis=alt.Axis(labelAngle=0)),
            y=alt.Y("valor:Q", title=None, axis=alt.Axis(format="~s")),
            color=alt.Color(
                "faixa:N",
                scale=alt.Scale(domain=list(FAIXAS_IDADE.values()), range=CORES_IDADE),
                legend=None,
            ),
            tooltip=[
                alt.Tooltip("faixa:N", title="Idade"),
                alt.Tooltip("texto:N", title="Valor a custo"),
                alt.Tooltip("detalhe:N", title="Itens"),
            ],
        ).properties(height=280)
        st.altair_chart(chart)

    with col_b:
        st.subheader("Valor a custo por categoria")
        df_itens = pd.DataFrame(itens)
        df_itens["valor_custo"] = df_itens["saldo"] * df_itens["preco_compra"]
        _grafico_horizontal(_top_agrupado(df_itens, "categoria", "valor_custo", 10), "Valor a custo")

    st.subheader("Maiores valores parados (>180 dias)")
    parados = sorted(
        (i for i in itens if i["classe_idade"] == "idade-vermelha"),
        key=lambda i: i["saldo"] * i["preco_compra"],
        reverse=True,
    )[:10]
    if not parados:
        st.caption("Nenhum produto com mais de 180 dias em estoque. 🎉")
        return
    st.dataframe(
        pd.DataFrame({
            "Cód. Produto": [i["codigo"] for i in parados],
            "Produto": [i["produto"] for i in parados],
            "Qtde": [_fmt_qtd(i["saldo"]) for i in parados],
            "Idade": [f"{i['idade']:.0f} dias" if i["idade"] is not None else "-" for i in parados],
            "Valor a custo": [_fmt_moeda(i["saldo"] * i["preco_compra"]) for i in parados],
        }),
        hide_index=True,
        use_container_width=True,
    )


# ----------------------------------------------------------------------
# Abas
# ----------------------------------------------------------------------
def _filtro_periodo():
    """Filtro de período (Compras e Vendas). Retorna (ini, fim) ou None se inválido."""
    hoje = date.today()
    opcoes = ["Todo período", "Ano atual", "Últimos 12 meses", "Últimos 90 dias", "Mês atual", "Personalizado"]

    col_p, col_a, col_b, col_btn = st.columns([2, 1.5, 1.5, 1.4], vertical_alignment="bottom")
    with col_p:
        escolha = st.selectbox("Período (compras e vendas)", opcoes, key="rep_periodo")

    ini, fim = None, hoje
    if escolha == "Todo período":
        fim = None
    elif escolha == "Ano atual":
        ini = date(hoje.year, 1, 1)
    elif escolha == "Últimos 12 meses":
        ini = hoje - timedelta(days=365)
    elif escolha == "Últimos 90 dias":
        ini = hoje - timedelta(days=90)
    elif escolha == "Mês atual":
        ini = date(hoje.year, hoje.month, 1)
    else:
        with col_a:
            ini = st.date_input("De", value=date(hoje.year, hoje.month, 1), key="rep_data_ini", format="DD/MM/YYYY")
        with col_b:
            fim = st.date_input("Até", value=hoje, key="rep_data_fim", format="DD/MM/YYYY")

    with col_btn:
        if st.button("🔄 Atualizar dados", use_container_width=True, key="rep_btn_atualizar"):
            st.cache_data.clear()
            st.rerun()

    if ini is not None and fim is not None and ini > fim:
        st.error("A data inicial não pode ser depois da data final.")
        return None
    return ini, fim


# Visões da aba "Reports - Geral": (rótulo, cor do destaque da aba ativa). As cores de
# Compras e Vendas acompanham as dos gráficos.
_VISOES_REPORT = [
    ("🛍️ Compras", COR_COMPRAS),
    ("💰 Vendas", COR_VENDAS),
    ("📦 Estoque", "#d97706"),
]


def _selecionar_visao(chave: str, valor: str):
    st.session_state[chave] = valor


def _css_navegacao(chave: str, visoes: list) -> str:
    """CSS da barra de abas: botões sem moldura, com sublinhado na aba ativa."""
    todos = f'[class*="st-key-{chave}_nav_"]'
    barra = f'div[data-testid="stHorizontalBlock"]:has({todos})'
    css = f"""
    {barra} {{
        flex-wrap: nowrap !important; gap: 0.25rem !important;
        border-bottom: 1px solid rgba(128,128,128,0.30); margin-bottom: 0.75rem;
    }}
    {barra} > div[data-testid="stColumn"] {{
        flex: 0 0 auto !important; width: auto !important; min-width: 0 !important;
    }}
    {todos} button {{
        background: transparent !important; color: inherit !important;
        border: none !important; border-bottom: 3px solid transparent !important;
        border-radius: 0 !important; box-shadow: none !important;
        padding: 0.55rem 1.15rem !important; opacity: 0.65;
    }}
    {todos} button p {{ font-size: 1.05rem !important; font-weight: 500 !important; color: inherit !important; }}
    {todos} button:hover {{ background: rgba(128,128,128,0.12) !important; opacity: 1; }}
    """
    for i, (_, cor) in enumerate(visoes):
        ativo = f'.st-key-{chave}_nav_{i} button[kind*="primary"]'
        css += (
            f"{ativo} {{ border-bottom-color: {cor} !important; opacity: 1; }}\n"
            f"{ativo} p {{ font-weight: 700 !important; }}\n"
        )
    return css


def _navegacao_visoes(visoes: list, chave: str) -> str:
    """Barra de navegação entre visões; devolve o rótulo da visão ativa.

    Continua sendo uma navegação por estado (e não st.tabs): só a visão aberta é
    executada e consulta o banco. Para voltar ao botão de rádio antigo, defina
    NAVEGACAO_VISOES = "radio".
    """
    rotulos = [r for r, _ in visoes]
    if st.session_state.get(chave) not in rotulos:
        st.session_state[chave] = rotulos[0]

    if NAVEGACAO_VISOES == "radio":
        return st.radio(
            "Visão", rotulos, horizontal=True, key=chave, label_visibility="collapsed"
        )

    atual = st.session_state[chave]
    st.markdown(f"<style>{_css_navegacao(chave, visoes)}</style>", unsafe_allow_html=True)
    for i, (coluna, rotulo) in enumerate(zip(st.columns(len(visoes)), rotulos)):
        with coluna:
            st.button(
                rotulo,
                key=f"{chave}_nav_{i}",
                type="primary" if rotulo == atual else "secondary",
                on_click=_selecionar_visao,
                args=(chave, rotulo),
            )
    return atual


def _aba_geral():
    periodo = _filtro_periodo()
    if periodo is None:
        return
    ini, fim = periodo

    visao = _navegacao_visoes(_VISOES_REPORT, "rep_visao")

    if visao == "🛍️ Compras":
        _visao_compras(ini, fim)
    elif visao == "💰 Vendas":
        _visao_vendas(ini, fim)
    else:
        _visao_estoque()


def _aba_bi():
    url_report = os.getenv("URL_REPORT")

    if not url_report:
        st.error(
            "⚠️ A URL do painel de relatórios não foi configurada no arquivo `.env`."
        )
        return

    st.markdown(
        """
        <style>
        div[data-testid="stLinkButton"] {
            width: 100% !important;
            display: flex !important;
            justify-content: center !important;
        }

        /* Estilo e largura fixa do botão */
        div[data-testid="stLinkButton"] a {
            background-color: #2e7d32 !important; /* Verde */
            color: #ffffff !important;
            border-color: #2e7d32 !important;
            width: 250px !important;               /* Largura do botão */
            text-align: center;
            font-weight: bold;
        }

        div[data-testid="stLinkButton"] a:hover {
            background-color: #1b5e20 !important; /* Verde escuro */
            border-color: #1b5e20 !important;
        }
        </style>
    """,
        unsafe_allow_html=True,
    )

    with st.container(border=True):
        st.markdown("### 📈 Painel Externo de Business Intelligence")
        st.write(
            "Acesse nossa plataforma externa para visualizar gráficos detalhados, "
            "métricas de vendas, análise de estoque e indicadores operacionais."
        )

        st.link_button(
            "🔗 Abrir Relatórios",
            url_report,
            use_container_width=False,
        )


def tela_report():
    st.header("📊 Relatórios e Analíticos")

    aba_geral, aba_bi = st.tabs(["Reports - Geral", "Reports - BI"])

    with aba_geral:
        _aba_geral()

    with aba_bi:
        _aba_bi()
