import streamlit as st
from datetime import date, datetime
from io import BytesIO

from db import supabase, buscar_todos
from componentes.paginacao import render_paginacao, get_itens_por_pagina, reset_paginacao
from componentes.campo_mascarado import campo_mascarado

_dialog = getattr(st, "dialog", None) or st.experimental_dialog


def _parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


def _fmt_date(value):
    parsed = _parse_date(value)
    return parsed.strftime("%d/%m/%Y") if parsed else "-"


def _parse_date_br(value):
    """Converte uma data digitada no componente para date.

    Campo vazio (inclusive máscara vazia) significa NULL. Aceita também
    separadores/espaços que o componente possa devolver durante o rerun.
    """
    if value is None:
        return None
    texto = str(value).strip()
    if not texto:
        return None

    # O componente de máscara trabalha com DD/MM/YYYY. Retiramos qualquer
    # separador para evitar falha quando o browser devolver espaços extras.
    digitos = "".join(ch for ch in texto if ch.isdigit())
    if not digitos:
        return None
    if len(digitos) != 8:
        return None

    try:
        return datetime.strptime(digitos, "%d%m%Y").date()
    except (ValueError, TypeError):
        return None


def _format_phone(value):
    digits = "".join(ch for ch in str(value or "") if ch.isdigit())[:11]
    if len(digits) <= 2:
        return f"({digits}" if digits else ""
    if len(digits) <= 7:
        return f"({digits[:2]}) {digits[2:]}"
    return f"({digits[:2]}) {digits[2:7]}-{digits[7:]}"



_COLUNAS_CLIENTE = "id, nome, telefone, canal_id, data_nascimento, data_casamento, detalhe_feira_id"

# rótulo -> (campo, decrescente?)
_ORDENACOES = {
    "Nome (A–Z)": ("nome", False),
    "Nome (Z–A)": ("nome", True),
    "Mais compras": ("qtde", True),
    "Menos compras": ("qtde", False),
    "Compra mais recente": ("ultima", True),
    "Compra mais antiga": ("ultima", False),
    "Maior valor comprado": ("valor", True),
    "Menor valor comprado": ("valor", False),
}
_ORDEM_PADRAO = "Nome (A–Z)"


def _limpar_filtros():
    st.session_state["clientes_filtro_nome"] = ""
    st.session_state["clientes_filtro_canal"] = "Todos os canais"
    st.session_state["clientes_ordem"] = _ORDEM_PADRAO
    reset_paginacao("clientes")


@_dialog("👤 Cliente", width="large")
def _dialog_cliente(cliente=None):
    editando = bool(cliente)
    titulo = "Editar cliente" if editando else "Adicionar cliente"
    st.markdown(f"### {titulo}")

    try:
        canais_resp = (
            supabase.table("canais_venda")
            .select("id, nome")
            .eq("ativo", True)
            .order("nome")
            .execute()
        )
        canais = canais_resp.data or []
    except Exception as err:
        st.error(f"Erro ao carregar canais: {err}")
        return

    try:
        feiras_resp = (
            supabase.table("detalhes_feira")
            .select("id, nome_feira")
            .eq("ativo", True)
            .order("nome_feira")
            .execute()
        )
        feiras = feiras_resp.data or []
    except Exception as err:
        st.error(f"Erro ao carregar feiras: {err}")
        return

    canal_ids = [c["id"] for c in canais]
    canal_labels = {c["id"]: c["nome"] for c in canais}
    feira_ids = [f["id"] for f in feiras]
    feira_labels = {f["id"]: f["nome_feira"] for f in feiras}
    cliente_id = cliente.get("id") if cliente else "novo"

    telefone_inicial = _format_phone(cliente.get("telefone") if cliente else "")

    col_nome, col_tel = st.columns([2.2, 1.3])
    with col_nome:
        nome = st.text_input(
            "Nome *",
            value=cliente.get("nome") or "" if cliente else "",
            key=f"cliente_nome_input_{cliente_id}",
        )
    with col_tel:
        telefone = campo_mascarado(
            "Telefone (opcional)",
            value=telefone_inicial,
            tipo="telefone",
            placeholder="(xx) xxxxx-xxxx",
            key=f"cliente_telefone_mask_{cliente_id}",
        )

    col_canal, col_feira, col_nasc, col_cas = st.columns([1.3, 1.5, 1.1, 1.1])
    with col_canal:
        atual_canal = cliente.get("canal_id") if cliente else None
        opcoes_canal = [None] + canal_ids
        canal_id = st.selectbox(
            "Canal *",
            options=opcoes_canal,
            index=opcoes_canal.index(atual_canal) if atual_canal in opcoes_canal else 0,
            format_func=lambda x: "-" if x is None else canal_labels[x],
            key=f"cliente_canal_input_{cliente_id}",
        )
    with col_feira:
        canal_e_feira = canal_id is not None and (canal_labels.get(canal_id) or "").strip().casefold() == "feira"
        feira_key = f"cliente_feira_input_{cliente_id}"
        atual_feira = cliente.get("detalhe_feira_id") if cliente else None
        opcoes_feira = [None] + feira_ids

        if not canal_e_feira:
            # Limpa de fato o valor do campo (não só na hora de salvar):
            # precisa ser feito ANTES de instanciar o widget abaixo, senão o
            # Streamlit mantém o que já estava selecionado em session_state
            # mesmo com o campo desabilitado.
            st.session_state[feira_key] = None

        detalhe_feira_id = st.selectbox(
            "Feira *" if canal_e_feira else "Feira (só p/ canal Feira)",
            options=opcoes_feira,
            index=opcoes_feira.index(atual_feira) if atual_feira in opcoes_feira else 0,
            format_func=lambda x: "-" if x is None else feira_labels[x],
            key=feira_key,
            disabled=not canal_e_feira,
            help=None if canal_e_feira else "Disponível apenas quando o canal selecionado é \"Feira\".",
        )
    with col_nasc:
        nascimento = campo_mascarado(
            "Data de nascimento (opcional)",
            value=_fmt_date(cliente.get("data_nascimento")) if cliente else "",
            tipo="data",
            placeholder="DD/MM/YYYY",
            key=f"cliente_nascimento_mask_{cliente_id}",
        )
    with col_cas:
        casamento = campo_mascarado(
            "Data de casamento (opcional)",
            value=_fmt_date(cliente.get("data_casamento")) if cliente else "",
            tipo="data",
            placeholder="DD/MM/YYYY",
            key=f"cliente_casamento_mask_{cliente_id}",
        )

    st.markdown("<div style='height: 18px'></div>", unsafe_allow_html=True)
    col_salvar, col_cancelar = st.columns(2)
    salvar = col_salvar.button(
        "💾 Salvar", type="primary", use_container_width=True, key=f"salvar_cliente_{cliente_id}"
    )
    cancelar = col_cancelar.button(
        "Cancelar", use_container_width=True, key=f"cancelar_cliente_{cliente_id}"
    )

    if cancelar:
        st.rerun()

    if not salvar:
        return

    nome = nome.strip()
    if not nome:
        st.error("Informe o nome do cliente.")
        return

    if canal_id is None:
        st.error("Selecione o canal.")
        return

    if canal_e_feira and detalhe_feira_id is None:
        st.error("Selecione a feira: é obrigatória quando o canal é \"Feira\".")
        return

    # Todos os campos, exceto nome, canal (e a feira quando o canal é
    # "Feira"), são opcionais. Normalizamos valores vazios/None para
    # permitir salvar NULL no banco.
    telefone = str(telefone or "").strip()
    nascimento = str(nascimento or "").strip()
    casamento = str(casamento or "").strip()

    # Campo de casamento/nascimento é opcional. Qualquer máscara vazia
    # (sem dígitos) deve ser tratada como NULL.
    if not any(ch.isdigit() for ch in nascimento):
        nascimento = ""
    if not any(ch.isdigit() for ch in casamento):
        casamento = ""

    nascimento_data = _parse_date_br(nascimento)
    casamento_data = _parse_date_br(casamento)

    if nascimento and not nascimento_data:
        st.error("Data de nascimento inválida. Use o formato DD/MM/YYYY.")
        return
    if casamento and not casamento_data:
        st.error("Data de casamento inválida. Use o formato DD/MM/YYYY.")
        return

    dados = {
        "nome": nome,
        "telefone": _format_phone(telefone) or None,
        "canal_id": canal_id if canal_id is not None else None,
        "detalhe_feira_id": detalhe_feira_id if detalhe_feira_id is not None else None,
        "data_nascimento": nascimento_data.isoformat() if nascimento_data else None,
        "data_casamento": casamento_data.isoformat() if casamento_data else None,
    }

    try:
        if editando:
            supabase.table("clientes").update(dados).eq("id", cliente["id"]).execute()
            st.toast("Cliente atualizado com sucesso!", icon="✅")
        else:
            supabase.table("clientes").insert(dados).execute()
            st.toast("Cliente cadastrado com sucesso!", icon="✅")
        st.rerun()
    except Exception as err:
        mensagem = str(err)
        # Detecta a violação da constraint de unicidade (nome + canal) e
        # mostra uma mensagem compreensível, em vez do erro cru do Postgres.
        # Cobre tanto o nome novo da constraint (uq_clientes_nome_canal)
        # quanto o antigo (uq_clientes_nome_lower), caso a migração ainda
        # não tenha sido rodada nesse banco.
        if "duplicate key value violates unique constraint" in mensagem and (
            "uq_clientes_nome_canal" in mensagem or "uq_clientes_nome_lower" in mensagem
        ):
            nome_canal = canal_labels.get(canal_id, "selecionado")
            st.error(
                f"Já existe um cliente chamado \"{nome}\" cadastrado no canal "
                f"\"{nome_canal}\". Se for a mesma pessoa, use o botão 🔗 Unificar "
                "na tela de clientes; se for outra pessoa, ajuste o nome ou "
                "confirme se o canal está correto."
            )
        else:
            st.error(f"Erro ao salvar cliente: {err}")


@_dialog("🔗 Unificar clientes", width="large")
def _dialog_unificar_clientes():
    """Migra as vendas de dois ou mais clientes "duplicados" para um único
    cliente "principal" e exclui os duplicados — pensado para o caso de
    cadastros repetidos por não existir CPF/RG como identificador único.

    Sempre que possível, a migração das vendas é feita por cliente_id
    (vínculo real). Para vendas antigas que ainda não tenham cliente_id
    preenchido, usamos a mesma comparação por nome (sem espaços nas pontas,
    sem diferenciar maiúsculas/minúsculas) já usada em _buscar_resumo_compras.
    """
    st.caption(
        "Selecione dois ou mais clientes duplicados **do mesmo canal**. As "
        "vendas de todos eles serão reatribuídas ao cliente principal "
        "escolhido, e os demais cadastros serão excluídos. Clientes de "
        "canais diferentes nunca são unificados entre si."
    )

    try:
        canais_resp = supabase.table("canais_venda").select("id, nome").order("nome").execute()
        canais_todos = canais_resp.data or []
    except Exception as err:
        st.error(f"Erro ao carregar canais: {err}")
        return
    canal_labels_todos = {c["id"]: c["nome"] for c in canais_todos}

    def _rotulo_canal(canal_id):
        return canal_labels_todos.get(canal_id, "Sem canal") if canal_id is not None else "Sem canal"

    try:
        todos_resp = (
            supabase.table("clientes")
            .select("id, nome, telefone, canal_id, data_nascimento, data_casamento, detalhe_feira_id")
            .order("nome")
            .execute()
        )
        todos_clientes = todos_resp.data or []
    except Exception as err:
        st.error(f"Erro ao carregar clientes: {err}")
        return

    if len(todos_clientes) < 2:
        st.info("É preciso ter pelo menos dois clientes cadastrados para unificar.")
        return

    # A unificação só é permitida DENTRO de um único canal — em vez de
    # validar depois de selecionar, o próprio formulário só oferece
    # candidatos do canal escolhido, tornando impossível misturar canais.
    canais_com_cliente = sorted(
        {cli.get("canal_id") for cli in todos_clientes},
        key=lambda cid: _rotulo_canal(cid).casefold(),
    )
    canal_escolhido = st.selectbox(
        "Canal",
        options=canais_com_cliente,
        format_func=_rotulo_canal,
        key="unificar_clientes_canal",
        help="Só é possível unificar clientes do mesmo canal.",
    )

    clientes_do_canal = [c for c in todos_clientes if c.get("canal_id") == canal_escolhido]

    if len(clientes_do_canal) < 2:
        st.info(f"Não há pelo menos dois clientes no canal \"{_rotulo_canal(canal_escolhido)}\" para unificar.")
        return

    clientes_por_id = {c["id"]: c for c in clientes_do_canal}
    rotulos = {
        c["id"]: c["nome"] + (f" — {c['telefone']}" if c.get("telefone") else "")
        for c in clientes_do_canal
    }

    selecionados = st.multiselect(
        "Clientes a unificar",
        options=[c["id"] for c in clientes_do_canal],
        format_func=lambda cid: rotulos[cid],
        key="unificar_clientes_selecionados",
    )

    if len(selecionados) < 2:
        st.info("Selecione pelo menos dois clientes para poder unificá-los.")
        return

    principal_id = st.radio(
        "Cliente principal (os demais serão migrados para este e excluídos)",
        options=selecionados,
        format_func=lambda cid: rotulos[cid],
        key="unificar_clientes_principal",
    )

    duplicados_ids = [cid for cid in selecionados if cid != principal_id]
    principal = clientes_por_id[principal_id]
    duplicados = [clientes_por_id[cid] for cid in duplicados_ids]
    nomes_duplicados_lower = {(d.get("nome") or "").strip().casefold() for d in duplicados}

    try:
        vendas_resp = supabase.table("vendas").select("id, cliente, cliente_id").execute()
        vendas_todas = vendas_resp.data or []
    except Exception as err:
        st.error(
            "Erro ao consultar vendas — verifique se a coluna `cliente_id` já "
            f"foi criada na tabela `vendas`. Detalhe: {err}"
        )
        return

    def _pertence_a_duplicado(venda: dict) -> bool:
        if venda.get("cliente_id") in duplicados_ids:
            return True
        if venda.get("cliente_id") is None:
            nome_venda = (venda.get("cliente") or "").strip().casefold()
            return nome_venda in nomes_duplicados_lower
        return False

    vendas_afetadas = [v for v in vendas_todas if _pertence_a_duplicado(v)]

    st.markdown(f"**Cliente principal:** {rotulos[principal_id]}")
    st.markdown("**Serão excluídos:** " + ", ".join(rotulos[cid] for cid in duplicados_ids))
    st.warning(
        f"{len(vendas_afetadas)} venda(s) terão o cliente reatribuído para "
        f"\"{principal.get('nome')}\". {len(duplicados_ids)} cadastro(s) de "
        "cliente serão excluídos permanentemente. Essa ação não pode ser desfeita."
    )

    confirmar = st.checkbox(
        "Confirmo que quero unificar esses clientes.",
        key="unificar_clientes_confirmar",
    )

    col_unificar, col_cancelar = st.columns(2)
    unificar = col_unificar.button(
        "🔗 Unificar clientes",
        type="primary",
        use_container_width=True,
        disabled=not confirmar,
    )
    cancelar = col_cancelar.button("Cancelar", use_container_width=True)

    if cancelar:
        st.rerun()

    if not unificar:
        return

    try:
        # 1. Reatribui as vendas afetadas (por cliente_id OU por nome em
        #    texto para vendas antigas), usando os IDs já calculados acima —
        #    evita depender de ILIKE no banco para casar nome com espaços/
        #    maiúsculas diferentes.
        ids_vendas_afetadas = [v["id"] for v in vendas_afetadas]
        if ids_vendas_afetadas:
            supabase.table("vendas").update({
                "cliente_id": principal_id,
                "cliente": principal.get("nome"),
            }).in_("id", ids_vendas_afetadas).execute()

        # 2. Completa dados do principal com o que os duplicados tiverem e
        #    o principal ainda não tiver preenchido.
        atualizacoes_principal = {}
        for campo in ("telefone", "canal_id", "detalhe_feira_id", "data_nascimento", "data_casamento"):
            if not principal.get(campo):
                for dup in duplicados:
                    if dup.get(campo):
                        atualizacoes_principal[campo] = dup[campo]
                        break
        if atualizacoes_principal:
            supabase.table("clientes").update(atualizacoes_principal).eq("id", principal_id).execute()

        # 3. Remove os cadastros duplicados.
        supabase.table("clientes").delete().in_("id", duplicados_ids).execute()

        st.toast("Clientes unificados com sucesso!", icon="✅")
        st.rerun()
    except Exception as err:
        st.error(f"Erro ao unificar clientes: {err}")


def _fmt_moeda(valor) -> str:
    try:
        v = float(valor or 0)
        return (
            f"R$ {v:,.2f}"
            .replace(",", "X")
            .replace(".", ",")
            .replace("X", ".")
        )
    except Exception:
        return "R$ 0,00"


def _buscar_vendas_do_cliente(cliente: dict) -> list[dict]:
    """Vendas de um cliente.

    Une os dois vínculos existentes no sistema, sem duplicar:
      1. vendas com `cliente_id` = id do cliente (vínculo real);
      2. vendas antigas/manuais sem `cliente_id`, casadas pelo nome (sem
         espaços nas pontas e sem diferenciar maiúsculas/minúsculas) — a
         mesma regra usada em _buscar_resumo_compras e na unificação.
    """
    colunas = (
        "id, data_venda, cliente, cliente_id, quantidade, valor_final, "
        "produtos(codigo_interno, descricao), canais_venda(nome), "
        "status_venda(nome), formas_pagamento(descricao)"
    )

    por_id = buscar_todos(
        lambda: supabase.table("vendas")
        .select(colunas)
        .eq("cliente_id", cliente["id"])
        .order("id")
    )

    vendas = {v["id"]: v for v in por_id}

    nome = (cliente.get("nome") or "").strip()
    if nome:
        # Escapa curingas do LIKE para o nome ser tratado literalmente. O
        # ilike só pré-filtra; a comparação exata é feita em Python abaixo.
        nome_like = nome.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        candidatas = buscar_todos(
            lambda: supabase.table("vendas")
            .select(colunas)
            .is_("cliente_id", "null")
            .ilike("cliente", f"%{nome_like}%")
            .order("id")
        )
        alvo = nome.casefold()
        for v in candidatas:
            if (v.get("cliente") or "").strip().casefold() == alvo:
                vendas.setdefault(v["id"], v)

    return sorted(
        vendas.values(),
        key=lambda v: (str(v.get("data_venda") or ""), v["id"]),
        reverse=True,
    )


@_dialog("🛒 Vendas do cliente", width="large")
def _dialog_vendas_cliente(cliente: dict, canal_labels: dict):
    st.markdown(f"### {cliente.get('nome') or '-'}")
    st.caption(
        f"Canal: {canal_labels.get(cliente.get('canal_id'), '-')}"
        + (f" · Telefone: {cliente['telefone']}" if cliente.get("telefone") else "")
    )

    try:
        vendas = _buscar_vendas_do_cliente(cliente)
    except Exception as err:
        st.error(f"Erro ao consultar as vendas do cliente: {err}")
        return

    if not vendas:
        st.info("Este cliente ainda não possui vendas.")
        return

    valor_total = sum(float(v.get("valor_final") or 0) for v in vendas)
    ultima = max((_parse_date(v.get("data_venda")) for v in vendas if v.get("data_venda")), default=None)

    col1, col2, col3 = st.columns(3)
    col1.metric("Itens vendidos", len(vendas))
    col2.metric("Valor total", _fmt_moeda(valor_total))
    col3.metric("Última compra", _fmt_date(ultima))

    st.dataframe(
        [
            {
                "Venda": v["id"],
                "Data": _fmt_date(v.get("data_venda")),
                "Produto": (v.get("produtos") or {}).get("descricao") or "-",
                "Código": (v.get("produtos") or {}).get("codigo_interno") or "-",
                "Qtd": v.get("quantidade"),
                "Valor": _fmt_moeda(v.get("valor_final")),
                "Canal": (v.get("canais_venda") or {}).get("nome") or "-",
                "Pagamento": (v.get("formas_pagamento") or {}).get("descricao") or "-",
                "Status": (v.get("status_venda") or {}).get("nome") or "-",
            }
            for v in vendas
        ],
        use_container_width=True,
        hide_index=True,
    )


def _buscar_resumo_compras() -> dict:
    """Resumo de compras por cliente: quantidade, última compra e valor.

    Segue a mesma regra de vínculo da janela "Vendas do cliente": a venda
    pertence ao cliente pelo `cliente_id` quando ele existe; as vendas sem
    `cliente_id` (antigas/manuais) são casadas pelo nome, sem espaços nas
    pontas e sem diferenciar maiúsculas/minúsculas.
    """
    resumo = {"por_id": {}, "por_nome": {}}
    try:
        vendas = buscar_todos(
            lambda: supabase.table("vendas")
            .select("id, cliente, cliente_id, data_venda, valor_final")
            .order("id")
        )
    except Exception as err:
        st.warning(f"Não foi possível carregar o resumo de compras dos clientes: {err}")
        return resumo

    for venda in vendas:
        if venda.get("cliente_id") is not None:
            bucket, chave = resumo["por_id"], venda["cliente_id"]
        else:
            nome = (venda.get("cliente") or "").strip().casefold()
            if not nome:
                continue
            bucket, chave = resumo["por_nome"], nome

        item = bucket.setdefault(chave, {"qtde": 0, "ultima": None, "valor": 0.0})
        item["qtde"] += 1
        item["valor"] += float(venda.get("valor_final") or 0)
        data = _parse_date(venda.get("data_venda"))
        if data and (item["ultima"] is None or data > item["ultima"]):
            item["ultima"] = data
    return resumo


def _stats_cliente(resumo: dict, cliente: dict) -> dict:
    """Soma o que está vinculado por `cliente_id` com o que casa pelo nome."""
    itens = (
        resumo["por_id"].get(cliente["id"]),
        resumo["por_nome"].get((cliente.get("nome") or "").strip().casefold()),
    )
    qtde, valor, ultima = 0, 0.0, None
    for item in itens:
        if not item:
            continue
        qtde += item["qtde"]
        valor += item["valor"]
        if item["ultima"] and (ultima is None or item["ultima"] > ultima):
            ultima = item["ultima"]
    return {"qtde": qtde, "valor": valor, "ultima": ultima}


def _query_clientes(filtro_nome: str, canal_id, count: bool = False):
    if count:
        query = supabase.table("clientes").select(_COLUNAS_CLIENTE, count="exact")
    else:
        query = supabase.table("clientes").select(_COLUNAS_CLIENTE)
    if filtro_nome.strip():
        query = query.ilike("nome", f"%{filtro_nome.strip()}%")
    if canal_id is not None:
        query = query.eq("canal_id", canal_id)
    return query


def _buscar_clientes_filtrados(filtro_nome: str, canal_id, ordem: str, resumo: dict) -> list[dict]:
    """TODOS os clientes dos filtros, já na ordem escolhida (sem paginar)."""
    campo, decrescente = _ORDENACOES.get(ordem, _ORDENACOES[_ORDEM_PADRAO])

    clientes = buscar_todos(
        lambda: _query_clientes(filtro_nome, canal_id)
        .order("nome", desc=(campo == "nome" and decrescente))
        .order("id")
    )
    if campo == "nome":
        return clientes

    # A lista parte ordenada por nome; o sort do Python é estável, então
    # clientes empatados na métrica continuam em ordem alfabética.
    if campo == "ultima":
        # Quem nunca comprou fica sempre no fim, em qualquer direção.
        com_compra = [c for c in clientes if _stats_cliente(resumo, c)["ultima"]]
        sem_compra = [c for c in clientes if not _stats_cliente(resumo, c)["ultima"]]
        com_compra.sort(key=lambda c: _stats_cliente(resumo, c)["ultima"], reverse=decrescente)
        return com_compra + sem_compra

    clientes.sort(key=lambda c: _stats_cliente(resumo, c)[campo], reverse=decrescente)
    return clientes


def _gerar_excel_clientes(linhas: list[dict]) -> bytes:
    """Planilha de clientes (uma linha por cliente) com valores numéricos
    reais, filtro, cabeçalho fixo e linha de total."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    cabecalhos = [
        "Nome", "Telefone", "Canal", "Feira", "Nascimento", "Casamento",
        "Última compra", "Qtd compras", "Valor comprado",
    ]
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    fonte = Font(color="FFFFFF", bold=True)
    for col, titulo in enumerate(cabecalhos, 1):
        c = ws.cell(1, col, titulo)
        c.fill = fill
        c.font = fonte
        c.alignment = Alignment(horizontal="center", vertical="center")

    linha = 2
    for item in linhas:
        valores = [
            item["nome"], item["telefone"], item["canal"], item["feira"],
            item["nascimento"], item["casamento"], item["ultima"],
            item["qtde"], item["valor"],
        ]
        for col, valor in enumerate(valores, 1):
            ws.cell(linha, col, valor)
        ws.cell(linha, 2).number_format = "@"  # telefone como texto
        for col in (5, 6, 7):
            ws.cell(linha, col).number_format = "DD/MM/YYYY"
            ws.cell(linha, col).alignment = Alignment(horizontal="center")
        ws.cell(linha, 8).number_format = "0"
        ws.cell(linha, 9).number_format = "R$ #,##0.00"
        linha += 1

    ultima_dados = linha - 1

    # SUBTOTAL: o total acompanha os filtros aplicados na própria planilha.
    linha_total = linha
    ws.cell(linha_total, 7, "TOTAL").font = Font(bold=True)
    ws.cell(linha_total, 7).alignment = Alignment(horizontal="right")
    borda = Border(top=Side(style="thin"))
    for col in range(1, len(cabecalhos) + 1):
        ws.cell(linha_total, col).border = borda
    for col in (8, 9):
        letra = get_column_letter(col)
        c = ws.cell(linha_total, col, f"=SUBTOTAL(109,{letra}2:{letra}{ultima_dados})")
        c.font = Font(bold=True)
        c.number_format = "0" if col == 8 else "R$ #,##0.00"

    larguras = [34, 18, 20, 24, 14, 14, 16, 13, 18]
    for i, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecalhos))}{ultima_dados}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _nome_arquivo_clientes(canal: str) -> str:
    """clientes_<canal>_<data>.xlsx, sem acentos nem caracteres inválidos."""
    import re
    import unicodedata

    sem_acento = unicodedata.normalize("NFKD", canal).encode("ascii", "ignore").decode()
    canal_limpo = re.sub(r"[^A-Za-z0-9]+", "-", sem_acento).strip("-") or "todos"
    return f"clientes_{canal_limpo}_{date.today():%Y%m%d}.xlsx"


@_dialog("📥 Exportar clientes", width="small")
def _dialog_exportar_clientes(filtro_nome: str, canal_id, canal_rotulo: str, ordem: str):
    """Busca TODOS os clientes dos filtros atuais (não só a página visível),
    na ordenação escolhida na tela, e oferece o download em Excel."""
    st.caption(
        f"Canal: **{canal_rotulo}**"
        + (f" · Nome contém: **{filtro_nome.strip()}**" if filtro_nome.strip() else "")
        + f" · Ordenação: **{ordem}**"
    )

    try:
        with st.spinner("Gerando planilha..."):
            resumo = _buscar_resumo_compras()
            clientes = _buscar_clientes_filtrados(filtro_nome, canal_id, ordem, resumo)

            canais = {
                c["id"]: c["nome"]
                for c in buscar_todos(
                    lambda: supabase.table("canais_venda").select("id, nome").order("id")
                )
            }
            feiras = {
                f["id"]: f["nome_feira"]
                for f in buscar_todos(
                    lambda: supabase.table("detalhes_feira").select("id, nome_feira").order("id")
                )
            }

            linhas = []
            for cliente in clientes:
                stats = _stats_cliente(resumo, cliente)
                linhas.append({
                    "nome": cliente.get("nome") or "",
                    "telefone": cliente.get("telefone") or "",
                    "canal": canais.get(cliente.get("canal_id"), ""),
                    "feira": feiras.get(cliente.get("detalhe_feira_id"), ""),
                    "nascimento": _parse_date(cliente.get("data_nascimento")),
                    "casamento": _parse_date(cliente.get("data_casamento")),
                    "ultima": stats["ultima"],
                    "qtde": stats["qtde"],
                    "valor": stats["valor"],
                })
            dados = _gerar_excel_clientes(linhas) if linhas else None
    except Exception as err:
        st.error(f"Não foi possível gerar a planilha: {err}")
        return

    if not linhas:
        st.info("Nenhum cliente para exportar com os filtros atuais.")
        return

    st.success(f"{len(linhas)} cliente(s) prontos para exportar.")
    st.download_button(
        "⬇️ Baixar Excel",
        data=dados,
        file_name=_nome_arquivo_clientes(canal_rotulo),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="clientes_download_excel",
    )


def _secao_clientes():
    try:
        canais_resp = (
            supabase.table("canais_venda")
            .select("id, nome")
            .eq("ativo", True)
            .order("nome")
            .execute()
        )
        canais = canais_resp.data or []
        canal_labels = {c["id"]: c["nome"] for c in canais}

        col_nome, col_canal, col_limpar, col_unificar, col_exportar, col_adicionar = st.columns(
            [2.0, 1.5, 1.0, 1.2, 1.2, 1.5]
        )

        with col_nome:
            filtro_nome = st.text_input(
                "Nome", key="clientes_filtro_nome", placeholder="Digite para buscar...", label_visibility="collapsed"
            )
        with col_canal:
            opcoes_canal = ["Todos os canais"] + [c["nome"] for c in canais]
            filtro_canal = st.selectbox(
                "Canal", options=opcoes_canal, key="clientes_filtro_canal", label_visibility="collapsed"
            )
        with col_limpar:
            st.button("Limpar filtros", use_container_width=True, on_click=_limpar_filtros)

        canal_selecionado = None
        if filtro_canal != "Todos os canais":
            canal_selecionado = next((c["id"] for c in canais if c["nome"] == filtro_canal), None)

        with col_unificar:
            if st.button("🔗 Unificar", use_container_width=True, help="Unificar clientes duplicados"):
                _dialog_unificar_clientes()
        with col_exportar:
            exportar = st.button(
                "📥 Exportar", use_container_width=True,
                help="Baixar os clientes dos filtros atuais em Excel",
            )
        with col_adicionar:
            if st.button("➕ Adicionar cliente", type="secondary", use_container_width=True):
                _dialog_cliente()

        col_ordem, _ = st.columns([2.0, 5.0])
        with col_ordem:
            ordem = st.selectbox(
                "Ordenar por",
                options=list(_ORDENACOES.keys()),
                key="clientes_ordem",
            )

        if exportar:
            _dialog_exportar_clientes(filtro_nome, canal_selecionado, filtro_canal, ordem)

        chave_filtro = f"{filtro_nome.strip().casefold()}|{filtro_canal}|{ordem}"
        if st.session_state.get("clientes_chave_filtro") != chave_filtro:
            st.session_state["clientes_chave_filtro"] = chave_filtro
            reset_paginacao("clientes")

        resumo = _buscar_resumo_compras()
        campo_ordem, decrescente = _ORDENACOES[ordem]

        if campo_ordem == "nome":
            # Ordenação por nome: paginação feita direto no banco.
            total_resp = (
                _query_clientes(filtro_nome, canal_selecionado, count=True)
                .order("nome", desc=decrescente).order("id")
                .range(0, 0).execute()
            )
            total = total_resp.count if total_resp.count is not None else 0

            def _carregar_pagina(offset, limite):
                resp = (
                    _query_clientes(filtro_nome, canal_selecionado)
                    .order("nome", desc=decrescente).order("id")
                    .range(offset, offset + limite - 1).execute()
                )
                return resp.data or []
        else:
            # Quantidade/última compra/valor dependem das vendas, então a
            # lista toda é ordenada em memória e só então paginada.
            todos = _buscar_clientes_filtrados(filtro_nome, canal_selecionado, ordem, resumo)
            total = len(todos)

            def _carregar_pagina(offset, limite):
                return todos[offset:offset + limite]

        st.header(f"👤 Clientes ({total} Clientes)")
        itens_por_pagina = get_itens_por_pagina("clientes")
        pagina = render_paginacao(
            "clientes", total, itens_por_pagina,
            mostrar_contagem_superior=True, mostrar_contagem_inferior=False, permitir_seletor=True
        )
        offset = (pagina - 1) * get_itens_por_pagina("clientes")

        clientes = _carregar_pagina(offset, get_itens_por_pagina("clientes"))

        if not clientes:
            st.info("Nenhum cliente encontrado para os filtros informados.")
            return

        larguras = [2.0, 1.2, 1.0, 1.0, 1.0, 1.1, 0.8, 1.2, 1.1]
        with st.container(border=True):
            cab = st.columns(larguras)
            for col, texto in zip(
                cab,
                ["Nome", "Telefone", "Canal", "Nascimento", "Casamento", "Última compra", "Compras", "Valor comprado", "Ações"],
            ):
                col.markdown(f"**{texto}**")
            st.divider()

            for cliente in clientes:
                stats = _stats_cliente(resumo, cliente)
                cols = st.columns(larguras)
                cols[0].write(cliente.get("nome") or "-")
                cols[1].write(cliente.get("telefone") or "-")
                cols[2].write(canal_labels.get(cliente.get("canal_id"), "-"))
                cols[3].write(_fmt_date(cliente.get("data_nascimento")))
                cols[4].write(_fmt_date(cliente.get("data_casamento")))
                cols[5].write(_fmt_date(stats["ultima"]))
                cols[6].write(str(stats["qtde"]))
                cols[7].write(_fmt_moeda(stats["valor"]))
                col_vendas, col_editar = cols[8].columns(2)
                if col_vendas.button(
                    "🛒", key=f"vendas_cliente_{cliente['id']}", help="Ver vendas do cliente"
                ):
                    _dialog_vendas_cliente(cliente, canal_labels)
                if col_editar.button(
                    "✏️", key=f"editar_cliente_{cliente['id']}", help="Editar cliente"
                ):
                    _dialog_cliente(cliente)

        render_paginacao(
            "clientes", total, get_itens_por_pagina("clientes"),
            mostrar_contagem_superior=False, mostrar_contagem_inferior=True, permitir_seletor=True
        )

    except Exception as err:
        st.error(f"Erro ao consultar clientes: {err}")


def tela_clientes():
    _secao_clientes()
